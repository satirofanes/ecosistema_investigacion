import sys
import os
import json
from PyQt6.QtWidgets import (QApplication, QWidget, QVBoxLayout, QHBoxLayout, 
                             QPushButton, QSplitter, QListWidget, QLabel, 
                             QTreeWidget, QTreeWidgetItem, QInputDialog, QMessageBox, 
                             QListWidgetItem, QAbstractItemView, QFileDialog, QComboBox, QDialog, QGraphicsView, QGraphicsScene)
from PyQt6.QtCore import Qt
from PyQt6.QtGui import QTextDocument, QPixmap, QColor
from sqlalchemy import create_engine, Column, Integer, String, Text
from sqlalchemy.orm import declarative_base, sessionmaker
import docx
from docx.shared import RGBColor

# --- CONFIGURACIÓN DE BD ---
Base = declarative_base()

class NotaPersonal(Base):
    __tablename__ = 'notas_personales'
    id = Column(Integer, primary_key=True)
    titulo = Column(String(150))
    desarrollo = Column(Text)
    hashtags = Column(String(200), default="")

engine = create_engine('sqlite:///investigacion_productiva.db')
Base.metadata.create_all(engine)
Session = sessionmaker(bind=engine)

# --- VISOR DE GUÍA VISUAL (Ventana Flotante) ---
class VisorGuia(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Guía Visual del Lienzo")
        self.resize(800, 600)
        layout = QVBoxLayout(self)

        btn_cargar = QPushButton("🖼️ Cargar Imagen PNG del Lienzo")
        btn_cargar.setStyleSheet("background-color: #3b82f6; color: white; padding: 6px;")
        btn_cargar.clicked.connect(self.cargar_imagen)

        self.escena = QGraphicsScene()
        self.visor = QGraphicsView(self.escena)
        self.visor.setDragMode(QGraphicsView.DragMode.ScrollHandDrag)
        self.visor.setStyleSheet("background-color: #1e293b;")

        layout.addWidget(btn_cargar)
        layout.addWidget(self.visor)

    def cargar_imagen(self):
        ruta, _ = QFileDialog.getOpenFileName(self, "Seleccionar Imagen del Lienzo", "", "Imágenes (*.png *.jpg *.jpeg)")
        if ruta:
            pixmap = QPixmap(ruta)
            self.escena.clear()
            item = self.escena.addPixmap(pixmap)
            self.visor.fitInView(item, Qt.AspectRatioMode.KeepAspectRatio)

    def wheelEvent(self, event):
        if event.modifiers() == Qt.KeyboardModifier.ControlModifier:
            zoom = 1.15 if event.angleDelta().y() > 0 else 1 / 1.15
            self.visor.scale(zoom, zoom)
        else:
            super().wheelEvent(event)

# --- MÓDULO VISUAL PRINCIPAL ---
class ModuloComposicion(QWidget):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Módulo 4: Integración, Organización y Delineados")
        self.resize(1200, 750)
        
        self.doc_limpiador = QTextDocument() 
        layout_principal = QVBoxLayout(self)
        
        # --- BARRA SUPERIOR ---
        layout_controles = QHBoxLayout()
        
        self.btn_refrescar = QPushButton("🔄 Refrescar Notas")
        self.btn_refrescar.clicked.connect(self.cargar_notas)

        self.btn_ver_lienzo = QPushButton("👁️ Ver Lienzo")
        self.btn_ver_lienzo.clicked.connect(self.abrir_guia)
        
        self.btn_exportar = QPushButton("📄 Exportar Estructura a Word (.docx)")
        self.btn_exportar.setStyleSheet("background-color: #2b579a; color: white; font-weight: bold; padding: 8px;")
        self.btn_exportar.clicked.connect(self.exportar_a_word)
        
        layout_controles.addWidget(self.btn_refrescar)
        layout_controles.addWidget(self.btn_ver_lienzo)
        layout_controles.addStretch()
        layout_controles.addWidget(self.btn_exportar)
        
        # --- ÁREA CENTRAL ---
        splitter = QSplitter(Qt.Orientation.Horizontal)
        
        # Panel Izquierdo
        panel_izq = QWidget()
        layout_izq = QVBoxLayout(panel_izq)
        layout_izq.setContentsMargins(0, 0, 10, 0)
        
        lbl_notas = QLabel("<b>1. Repositorio de Ideas</b><br><i>(Doble clic para enviar a la estructura ➔)</i>")
        self.lista_notas = QListWidget()
        self.lista_notas.itemDoubleClicked.connect(self.enviar_nota_al_arbol)
        self.lista_notas.setWordWrap(True)
        
        layout_izq.addWidget(lbl_notas)
        layout_izq.addWidget(self.lista_notas)
        
        # Panel Derecho
        panel_der = QWidget()
        layout_der = QVBoxLayout(panel_der)
        layout_der.setContentsMargins(10, 0, 0, 0)
        
        lbl_arbol = QLabel("<b>2. Esqueleto del Delineado</b><br><i>(Arrastra los elementos para anidarlos u ordenarlos)</i>")
        layout_der.addWidget(lbl_arbol)

        # Controles del árbol
        layout_controles_arbol = QHBoxLayout()
        self.btn_add_seccion = QPushButton("📁 Nueva Sección")
        self.btn_add_seccion.clicked.connect(self.crear_seccion)
        
        # Opciones de Preposiciones
        self.combo_prep = QComboBox()
        self.combo_prep.addItems(["a", "ante", "con", "contra", "en", "entre", "hacia", "hasta", "para", "por", "según", "sin", "sobre", "tras"])
        self.btn_add_prep = QPushButton("🔗 Nexo")
        self.btn_add_prep.clicked.connect(self.crear_preposicion)

        self.btn_eliminar_nodo = QPushButton("🗑️ Quitar Nodo")
        self.btn_eliminar_nodo.clicked.connect(self.eliminar_nodo_arbol)
        
        layout_controles_arbol.addWidget(self.btn_add_seccion)
        layout_controles_arbol.addWidget(self.combo_prep)
        layout_controles_arbol.addWidget(self.btn_add_prep)
        layout_controles_arbol.addStretch()
        layout_controles_arbol.addWidget(self.btn_eliminar_nodo)
        
        self.arbol = QTreeWidget()
        self.arbol.setColumnCount(2)
        self.arbol.hideColumn(1) 
        self.arbol.setHeaderHidden(True)
        self.arbol.setDragDropMode(QAbstractItemView.DragDropMode.InternalMove)
        self.arbol.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.arbol.setStyleSheet("QTreeWidget { font-size: 14px; } QTreeWidget::item { padding: 6px; border-bottom: 1px solid #e2e8f0; }")
        
        layout_der.addLayout(layout_controles_arbol)
        layout_der.addWidget(self.arbol)
        
        splitter.addWidget(panel_izq)
        splitter.addWidget(panel_der)
        splitter.setSizes([450, 750])
        
        layout_principal.addLayout(layout_controles)
        layout_principal.addWidget(splitter)
        
        self.cargar_notas()

    def abrir_guia(self):
        visor = VisorGuia(self)
        visor.show()

    def cargar_notas(self):
        self.lista_notas.clear()
        session = Session()
        notas = session.query(NotaPersonal).all()
        
        for nota in notas:
            titulo = nota.titulo if nota.titulo else "Sin título"
            self.doc_limpiador.setHtml(nota.desarrollo or "")
            extracto = self.doc_limpiador.toPlainText().replace('\n', ' ')[:60]
            
            item = QListWidgetItem(f"📝 {titulo}\n{extracto}...")
            datos_json = json.dumps({
                "type": "nota",
                "titulo": titulo,
                "contenido_html": nota.desarrollo or "",
                "hashtags": nota.hashtags or ""
            })
            item.setData(Qt.ItemDataRole.UserRole, datos_json)
            self.lista_notas.addItem(item)
            
        session.close()

    def enviar_nota_al_arbol(self, item_lista):
        datos_json = item_lista.data(Qt.ItemDataRole.UserRole)
        datos = json.loads(datos_json)
        
        item_arbol = QTreeWidgetItem([f"📝 {datos['titulo']}", datos_json])
        nodo_seleccionado = self.arbol.currentItem()
        
        if nodo_seleccionado and json.loads(nodo_seleccionado.text(1))["type"] == "seccion":
            nodo_seleccionado.addChild(item_arbol)
            nodo_seleccionado.setExpanded(True)
        else:
            self.arbol.addTopLevelItem(item_arbol)

    def crear_seccion(self):
        nombre, ok = QInputDialog.getText(self, "Nueva Sección", "Nombre del Capítulo o Sección:")
        if ok and nombre.strip():
            datos_json = json.dumps({"type": "seccion", "titulo": nombre})
            item_arbol = QTreeWidgetItem([f"📁 {nombre}", datos_json])
            font = item_arbol.font(0)
            font.setBold(True)
            item_arbol.setFont(0, font)
            self.arbol.addTopLevelItem(item_arbol)

    def crear_preposicion(self):
        texto = self.combo_prep.currentText()
        datos_json = json.dumps({"type": "preposicion", "texto": texto})
        item_arbol = QTreeWidgetItem([f"🔗 {texto}", datos_json])
        item_arbol.setForeground(0, QColor("#8b5cf6")) 

        nodo_seleccionado = self.arbol.currentItem()
        if nodo_seleccionado and json.loads(nodo_seleccionado.text(1))["type"] == "seccion":
            nodo_seleccionado.addChild(item_arbol)
            nodo_seleccionado.setExpanded(True)
        else:
            self.arbol.addTopLevelItem(item_arbol)

    def eliminar_nodo_arbol(self):
        nodo_seleccionado = self.arbol.currentItem()
        if nodo_seleccionado:
            if nodo_seleccionado.parent():
                nodo_seleccionado.parent().removeChild(nodo_seleccionado)
            else:
                self.arbol.takeTopLevelItem(self.arbol.indexOfTopLevelItem(nodo_seleccionado))

    def exportar_a_word(self):
        try:
            if self.arbol.topLevelItemCount() == 0:
                QMessageBox.warning(self, "Aviso", "El delineado está vacío. Añade elementos primero.")
                return

            ruta_guardado, _ = QFileDialog.getSaveFileName(
                self, "Guardar Delineado en Word", "Borrador_Investigacion.docx", "Documento de Word (*.docx)",
                options=QFileDialog.Option.DontUseNativeDialog
            )
            if not ruta_guardado: return

            documento = docx.Document()
            documento.add_heading('Delineado de Investigación', 0)

            def procesar_nodo(nodo, nivel_actual):
                texto_json = nodo.text(1)
                if not texto_json: return
                datos = json.loads(texto_json)
                
                if datos["type"] == "seccion":
                    documento.add_heading(datos["titulo"], level=min(nivel_actual, 4))
                    for i in range(nodo.childCount()):
                        procesar_nodo(nodo.child(i), nivel_actual + 1)
                        
                elif datos["type"] == "nota":
                    self.doc_limpiador.setHtml(datos["contenido_html"])
                    documento.add_heading(datos["titulo"], level=min(nivel_actual, 4))
                    
                    if datos.get("hashtags"):
                        p_tags = documento.add_paragraph()
                        r = p_tags.add_run(f"{datos['hashtags']}")
                        r.italic = True
                        r.font.color.rgb = RGBColor(120, 120, 120)

                    documento.add_paragraph(self.doc_limpiador.toPlainText())
                    
                    for i in range(nodo.childCount()):
                        procesar_nodo(nodo.child(i), nivel_actual)
                        
                elif datos["type"] == "preposicion":
                    p = documento.add_paragraph()
                    r = p.add_run(f" — {datos['texto']} — ")
                    r.italic = True
                    r.bold = True
                    p.alignment = 1 
                    
                    for i in range(nodo.childCount()):
                        procesar_nodo(nodo.child(i), nivel_actual)

            for i in range(self.arbol.topLevelItemCount()):
                procesar_nodo(self.arbol.topLevelItem(i), 1)

            documento.save(ruta_guardado)
            os.startfile(os.path.normpath(ruta_guardado)) 
            
        except Exception as e:
            QMessageBox.critical(self, "Error Fatal", f"Ocurrió un fallo al generar el Word:\n\n{str(e)}")

if __name__ == "__main__":
    app = QApplication(sys.argv)
    app.setStyle("Fusion")
    ventana = ModuloComposicion()
    ventana.show()
    sys.exit(app.exec())
